"""Documents: extract text with page numbers, chunk it, embed chunks with a local text model.

Formats: PDF (PyMuPDF), DOCX (python-docx), TXT / Markdown. Scanned PDFs without a text layer are
not OCR'd (kept simple on purpose): they are indexed by title only.
"""

import io
import re
import threading
from dataclasses import dataclass

CHUNK_WORDS = 300
OVERLAP_WORDS = 50


@dataclass
class Chunk:
    page: int | None  # 1-based page (PDF) / estimated page (DOCX); None for plain text
    index: int
    text: str


# ---------------------------------------------------------------- extraction
def _clean(text: str) -> str:
    text = text.replace("­", "")  # soft hyphens
    text = re.sub(r"-\n(\w)", r"\1", text)  # re-join words hyphenated across lines
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _pdf(data: bytes) -> tuple[str | None, list[tuple[int | None, str]]]:
    import pymupdf

    with pymupdf.open(stream=data, filetype="pdf") as doc:
        title = (doc.metadata or {}).get("title") or None
        pages = [(i + 1, _clean(page.get_text("text"))) for i, page in enumerate(doc)]
    return title, pages


def _docx(data: bytes) -> tuple[str | None, list[tuple[int | None, str]]]:
    import docx

    d = docx.Document(io.BytesIO(data))
    title = (d.core_properties.title or "").strip() or None
    pages: list[tuple[int | None, str]] = []
    page, buf = 1, []
    for p in d.paragraphs:
        xml = p._p.xml
        # Word marks page boundaries with explicit breaks and (when saved by Word) rendered breaks.
        breaks = xml.count('w:type="page"') + xml.count("lastRenderedPageBreak")
        if breaks and buf:
            pages.append((page, _clean("\n".join(buf))))
            buf = []
        page += breaks
        if p.text.strip():
            if title is None and p.style is not None and p.style.name.lower().startswith(("title", "heading 1")):
                title = p.text.strip()
            buf.append(p.text)
    for table in d.tables:  # tables are separate from paragraphs in python-docx
        for row in table.rows:
            buf.append(" | ".join(c.text.strip() for c in row.cells if c.text.strip()))
    if buf:
        pages.append((page, _clean("\n".join(buf))))
    return title, pages


def _text(data: bytes, markdown: bool) -> tuple[str | None, list[tuple[int | None, str]]]:
    text = data.decode("utf-8", errors="replace")
    title = None
    if markdown:
        m = re.search(r"^#\s+(.+)$", text, re.M)
        title = m.group(1).strip() if m else None
    return title, [(None, _clean(text))]


def extract(data: bytes, filename: str) -> tuple[str, list[tuple[int | None, str]]]:
    """Returns (title, [(page, text), ...]). Raises ValueError for unsupported types."""
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "pdf" or data[:5] == b"%PDF-":
        title, pages = _pdf(data)
    elif ext == "docx":
        title, pages = _docx(data)
    elif ext in ("md", "markdown"):
        title, pages = _text(data, markdown=True)
    elif ext == "txt":
        title, pages = _text(data, markdown=False)
    else:
        raise ValueError(f"unsupported document type: .{ext}")
    fallback = re.sub(r"[_-]+", " ", filename.rsplit(".", 1)[0]).strip() or "Untitled"
    return (title or fallback)[:300], pages


# ---------------------------------------------------------------- chunking
def chunk(pages: list[tuple[int | None, str]], size: int = CHUNK_WORDS, overlap: int = OVERLAP_WORDS) -> list[Chunk]:
    """~size-word windows with overlap, never crossing a page (so every chunk has an exact page)."""
    out: list[Chunk] = []
    step = max(1, size - overlap)
    for page, text in pages:
        words = text.split()
        if not words:
            continue
        for start in range(0, len(words), step):
            piece = words[start : start + size]
            out.append(Chunk(page=page, index=len(out), text=" ".join(piece)))
            if start + size >= len(words):
                break
    return out


def snippet(text: str, query: str, width: int = 40) -> str:
    """A ~width-word window around the first query word found in the chunk."""
    words = text.split()
    terms = [t for t in re.findall(r"\w+", query.lower()) if len(t) > 2]
    hit = next((i for i, w in enumerate(words) if any(t in w.lower() for t in terms)), 0)
    start = max(0, hit - width // 3)
    piece = " ".join(words[start : start + width])
    return ("… " if start > 0 else "") + piece + (" …" if start + width < len(words) else "")


# ---------------------------------------------------------------- embeddings
class TextEmbedder:
    """multilingual-e5-small: small, CPU-friendly, multilingual. e5 expects 'query:' / 'passage:' prefixes."""

    def __init__(self, model_name: str, threads: int = 0):
        import os

        import torch
        from sentence_transformers import SentenceTransformer

        torch.set_num_threads(threads or os.cpu_count() or 1)
        self.name = model_name
        self.model = SentenceTransformer(model_name, device="cpu")
        self._lock = threading.Lock()

    def passages(self, texts: list[str], batch: int = 16) -> list[list[float]]:
        # The lock is taken per batch, not per document: a search query waits for one batch at most,
        # not for a whole 900-page book.
        out: list[list[float]] = []
        for i in range(0, len(texts), batch):
            with self._lock:
                vecs = self.model.encode([f"passage: {t}" for t in texts[i : i + batch]], normalize_embeddings=True)
            out.extend(v.tolist() for v in vecs)
        return out

    def query(self, text: str) -> list[float]:
        with self._lock:
            return self.model.encode([f"query: {text}"], normalize_embeddings=True)[0].tolist()
