import io

import pytest

from app import documents as d


def make_pdf(pages: list[str]) -> bytes:
    import pymupdf

    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        page.insert_textbox(pymupdf.Rect(50, 50, 550, 800), text, fontsize=11)
    doc.set_metadata({"title": "Database Notes"})
    data = doc.tobytes()
    doc.close()
    return data


def make_docx() -> bytes:
    import docx
    from docx.enum.text import WD_BREAK

    doc = docx.Document()
    doc.add_heading("Quarterly Report", level=1)
    doc.add_paragraph("Revenue grew in the first quarter thanks to new customers.")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph("The budget table shows marketing spend by region.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_pdf_pages_and_title():
    title, pages = d.extract(make_pdf(["Intro to relational databases.", "Normalization removes redundancy."]), "x.pdf")
    assert title == "Database Notes"
    assert [p for p, _ in pages] == [1, 2]
    assert "Normalization" in pages[1][1]


def test_docx_title_and_page_breaks():
    title, pages = d.extract(make_docx(), "report.docx")
    assert title == "Quarterly Report"
    assert pages[0][0] == 1 and "Revenue" in pages[0][1]
    assert pages[-1][0] == 2 and "budget table" in pages[-1][1]


def test_markdown_title_and_plain_text():
    title, pages = d.extract(b"# Caching guide\n\nUse a TTL.", "guide.md")
    assert title == "Caching guide" and pages == [(None, "# Caching guide\n\nUse a TTL.")]
    title, _ = d.extract(b"hello", "my_notes-v2.txt")
    assert title == "my notes v2"


def test_unsupported_type():
    with pytest.raises(ValueError):
        d.extract(b"PK...", "deck.pptx")


def test_chunking_overlaps_and_never_crosses_pages():
    words = " ".join(f"w{i}" for i in range(700))
    chunks = d.chunk([(1, words), (2, "short page")], size=300, overlap=50)
    assert [c.page for c in chunks] == [1, 1, 1, 2]
    assert chunks[0].text.split()[-50:] == chunks[1].text.split()[:50]  # 50-word overlap
    assert chunks[-1].text == "short page"


def test_snippet_centres_on_query_words():
    text = " ".join(["filler"] * 60 + ["normalization", "removes", "redundancy"] + ["more"] * 60)
    s = d.snippet(text, "what is normalization", width=20)
    assert "normalization" in s and s.startswith("…") and s.endswith("…")
